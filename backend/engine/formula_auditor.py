import re
from typing import Dict, List, Any, Optional, Tuple, Set
import openpyxl
from openpyxl.utils import coordinate_to_tuple, get_column_letter

class FormulaAuditor:
    """
    Audits formula consistency, detects formula pattern breaks (hardcoded overrides),
    identifies range boundary slippages (e.g. SUM omitting trailing rows),
    and warns about hidden row inclusion risks.
    """

    @classmethod
    def audit_workbook(cls, wb: openpyxl.Workbook, data_wb: Optional[openpyxl.Workbook] = None) -> Dict[str, Any]:
        sheet_audits = {}
        all_findings = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            data_ws = data_wb[sheet_name] if data_wb and sheet_name in data_wb.sheetnames else None
            audit = cls.audit_worksheet(ws, data_ws)
            sheet_audits[sheet_name] = audit
            all_findings.extend(audit["findings"])

        return {
            "sheets": sheet_audits,
            "findings": all_findings,
            "total_formula_issues": len(all_findings)
        }

    @classmethod
    def audit_worksheet(
        cls, 
        ws: openpyxl.worksheet.worksheet.Worksheet,
        data_ws: Optional[openpyxl.worksheet.worksheet.Worksheet] = None
    ) -> Dict[str, Any]:
        findings = []
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1

        if max_row < 2:
            return {"sheet_name": ws.title, "findings": [], "summary_cells": []}

        # Check for hidden rows
        hidden_rows = {
            r for r, dim in ws.row_dimensions.items()
            if dim.hidden and r <= max_row
        }

        # Group cells by column to detect formula pattern continuity
        col_cells: Dict[int, List[Tuple[int, Any, bool]]] = {c: [] for c in range(1, max_col + 1)}
        summary_cells = []

        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                val = cell.value
                is_formula = (
                    cell.data_type == 'f' or 
                    (isinstance(val, str) and val.startswith('='))
                )
                col_cells[c].append((r, val, is_formula))

                if is_formula:
                    formula_str = str(val).upper()
                    if any(fn in formula_str for fn in ["SUM", "AVERAGE", "COUNT", "MAX", "MIN", "SUBTOTAL"]):
                        summary_cells.append({
                            "coordinate": cell.coordinate,
                            "row": r,
                            "col": c,
                            "formula": str(val)
                        })

        # 1. Detect Formula Pattern Breaks (Hardcoded overrides in formula columns)
        for col_idx, cells in col_cells.items():
            cls._check_column_formula_continuity(ws, col_idx, cells, findings)

        # 2. Detect Range Boundary Slippage in Summary formulas
        for summary in summary_cells:
            cls._check_range_slippage(ws, summary, max_row, findings, data_ws)

        # 3. Detect Hidden Row inclusion risk
        for summary in summary_cells:
            cls._check_hidden_row_risk(ws, summary, hidden_rows, findings)

        return {
            "sheet_name": ws.title,
            "findings": findings,
            "summary_cells": summary_cells
        }

    @classmethod
    def _check_column_formula_continuity(
        cls, ws: openpyxl.worksheet.worksheet.Worksheet, col_idx: int, 
        cells: List[Tuple[int, Any, bool]], findings: List[Dict[str, Any]]
    ):
        """
        Detects if a column is predominantly calculated via formula, but some rows
        are hardcoded numbers or altered formulas.
        """
        # Exclude row 1 (assumed header)
        data_cells = [(r, val, is_form) for r, val, is_form in cells if r > 1 and val is not None]
        if len(data_cells) < 4:
            return

        formula_cells = [(r, val) for r, val, is_form in data_cells if is_form]
        hardcoded_cells = [(r, val) for r, val, is_form in data_cells if not is_form and isinstance(val, (int, float))]

        total_populated = len(data_cells)
        formula_ratio = len(formula_cells) / total_populated

        # If 70% or more are formulas, but some are hardcoded numbers
        if formula_ratio >= 0.70 and hardcoded_cells:
            col_letter = get_column_letter(col_idx)
            header_val = str(ws.cell(row=1, column=col_idx).value or f"Column {col_letter}")

            # Extract formula templates to see what the dominant pattern is
            patterns = {}
            for r, f_val in formula_cells:
                # Normalize row numbers: replace exact row with {r}
                norm = re.sub(rf"(?<=[A-Z]){r}\b", "{r}", str(f_val))
                patterns[norm] = patterns.get(norm, 0) + 1

            dominant_pattern = max(patterns.items(), key=lambda x: x[1])[0] if patterns else "formula"

            for r, hard_val in hardcoded_cells:
                # Don't flag the very last row if it's a total row
                if r == ws.max_row:
                    continue

                cell_ref = f"{col_letter}{r}"
                expected_formula = dominant_pattern.replace("{r}", str(r))

                findings.append({
                    "id": f"break_{cell_ref}",
                    "category": "formula_integrity",
                    "severity": "critical",
                    "type": "hardcoded_override",
                    "cell": cell_ref,
                    "row": r,
                    "column": col_letter,
                    "column_name": header_val,
                    "current_value": hard_val,
                    "proposed_value": expected_formula,
                    "message": f"Hardcoded value '{hard_val}' breaks formula pattern in column '{header_val}'. Contiguous rows calculate '{expected_formula}'.",
                    "impact_hint": "May result in stale calculations if input assumptions change.",
                    "can_autofix": True
                })

    @classmethod
    def _check_range_slippage(
        cls, ws: openpyxl.worksheet.worksheet.Worksheet, summary: Dict[str, Any],
        max_row: int, findings: List[Dict[str, Any]],
        data_ws: Optional[openpyxl.worksheet.worksheet.Worksheet] = None
    ):
        """
        Detects formulas like =SUM(D2:D45) where non-empty data rows extend to row 55.
        """
        formula = summary["formula"]
        # Match ranges like D2:D45 or $D$2:$D$45
        range_matches = re.findall(r"([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", formula.upper())

        for start_col, start_r, end_col, end_r in range_matches:
            if start_col == end_col:
                col_letter = start_col
                s_row = int(start_r)
                e_row = int(end_r)

                # Check if there is data between e_row + 1 and the summary cell row
                target_col_idx = openpyxl.utils.column_index_from_string(col_letter)
                summary_row = summary["row"]

                # Check rows between e_row + 1 and summary_row - 1
                omitted_populated_rows = []
                for check_r in range(e_row + 1, summary_row):
                    cell_val = ws.cell(row=check_r, column=target_col_idx).value
                    # If data_ws exists, get evaluated numeric value
                    num_val = data_ws.cell(row=check_r, column=target_col_idx).value if data_ws else cell_val
                    if cell_val is not None and str(cell_val).strip() != "":
                        omitted_populated_rows.append((check_r, num_val if isinstance(num_val, (int, float)) else 0.0))

                if omitted_populated_rows:
                    omitted_count = len(omitted_populated_rows)
                    omitted_sum = sum(v for _, v in omitted_populated_rows if isinstance(v, (int, float)))
                    proposed_formula = formula.replace(f"{col_letter}{e_row}", f"{col_letter}{omitted_populated_rows[-1][0]}")

                    findings.append({
                        "id": f"slippage_{summary['coordinate']}",
                        "category": "range_slippage",
                        "severity": "critical",
                        "type": "omitted_rows_in_sum",
                        "cell": summary["coordinate"],
                        "row": summary_row,
                        "column": get_column_letter(summary["col"]),
                        "column_name": str(ws.cell(row=1, column=summary["col"]).value or col_letter),
                        "current_value": formula,
                        "proposed_value": proposed_formula,
                        "omitted_rows": [r for r, _ in omitted_populated_rows],
                        "omitted_count": omitted_count,
                        "omitted_sum": omitted_sum,
                        "message": f"Summary formula '{formula}' at {summary['coordinate']} stops at row {e_row}, omitting {omitted_count} populated rows (rows {omitted_populated_rows[0][0]} to {omitted_populated_rows[-1][0]}).",
                        "impact_hint": f"Excludes approximately {omitted_sum:,.2f} from total calculation.",
                        "can_autofix": True
                    })

    @classmethod
    def _check_hidden_row_risk(
        cls, ws: openpyxl.worksheet.worksheet.Worksheet, summary: Dict[str, Any],
        hidden_rows: Set[int], findings: List[Dict[str, Any]]
    ):
        """
        Warns if standard =SUM() covers hidden rows that might have been filtered out.
        """
        if not hidden_rows:
            return

        formula = summary["formula"].upper()
        if "=SUM(" in formula and not any(fn in formula for fn in ["SUBTOTAL", "AGGREGATE"]):
            range_matches = re.findall(r"([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", formula)
            for start_col, start_r, end_col, end_r in range_matches:
                s_row, e_row = int(start_r), int(end_r)
                intersecting_hidden = [r for r in hidden_rows if s_row <= r <= e_row]
                if intersecting_hidden:
                    findings.append({
                        "id": f"hidden_{summary['coordinate']}",
                        "category": "hidden_cell_risk",
                        "severity": "medium",
                        "type": "sum_includes_hidden_rows",
                        "cell": summary["coordinate"],
                        "row": summary["row"],
                        "column": get_column_letter(summary["col"]),
                        "column_name": "Summary Calculation",
                        "current_value": summary["formula"],
                        "proposed_value": summary["formula"].replace("SUM(", "SUBTOTAL(109, "),
                        "message": f"Formula '{summary['formula']}' in {summary['coordinate']} uses =SUM() which blindly calculates values from {len(intersecting_hidden)} hidden/filtered rows.",
                        "impact_hint": "May double-count or include excluded line items in executive summaries.",
                        "can_autofix": True
                    })
