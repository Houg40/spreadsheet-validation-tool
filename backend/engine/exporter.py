import os
from datetime import datetime
from typing import Dict, List, Any
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

class SafeSpreadsheetExporter:
    """
    Applies staged accepted patches to spreadsheets non-destructively,
    adds cell comments with audit provenance, and appends a formal _Audit_Changelog worksheet.
    """

    @classmethod
    def export_cleaned_workbook(
        cls,
        source_path: str,
        output_path: str,
        accepted_patches: List[Dict[str, Any]],
        mode: str = "annotated" # 'clean' or 'annotated'
    ) -> str:
        wb = openpyxl.load_workbook(source_path, data_only=False)

        # Style definitions
        modified_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Soft green
        audit_header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        audit_header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

        audit_entries = []

        for patch in accepted_patches:
            sheet_name = patch.get("sheet_name") or wb.sheetnames[0]
            if sheet_name not in wb.sheetnames:
                continue

            ws = wb[sheet_name]
            coord = patch.get("cell")
            if not coord:
                continue

            orig_val = patch.get("current_value")
            new_val = patch.get("proposed_value")
            reason = patch.get("message", "Rule validation fix")

            cell = ws[coord]
            cell.value = new_val

            if mode == "annotated":
                cell.fill = modified_fill
                comment_text = (
                    f"Spreadsheet Quality Inspector Audit:\n"
                    f"• Modified: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
                    f"• Original Value: {orig_val}\n"
                    f"• Applied Fix: {new_val}\n"
                    f"• Rationale: {reason}"
                )
                cell.comment = Comment(comment_text, "Spreadsheet Inspector")

            audit_entries.append({
                "Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Sheet": sheet_name,
                "Cell": coord,
                "Field": patch.get("column_name", ""),
                "Original Value": str(orig_val),
                "Cleaned Value": str(new_val),
                "Category": patch.get("category", "General"),
                "Impact Statement": patch.get("impact", {}).get("executive_statement", reason),
                "Status": "Accepted"
            })

        if mode == "annotated":
            # Add or update _Audit_Changelog tab
            if "_Audit_Changelog" in wb.sheetnames:
                del wb["_Audit_Changelog"]

            log_ws = wb.create_sheet(title="_Audit_Changelog")
            headers = ["Timestamp", "Sheet", "Cell", "Field", "Original Value", "Cleaned Value", "Category", "Impact Statement", "Status"]
            log_ws.append(headers)

            for col_idx in range(1, len(headers) + 1):
                c = log_ws.cell(row=1, column=col_idx)
                c.fill = audit_header_fill
                c.font = audit_header_font

            for entry in audit_entries:
                log_ws.append([entry[h] for h in headers])

            # Auto-fit columns
            for col in log_ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                log_ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

        wb.save(output_path)
        return output_path
