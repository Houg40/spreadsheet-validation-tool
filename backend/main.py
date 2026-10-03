import os
import shutil
import uuid
from typing import Dict, List, Any, Optional
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
import openpyxl
import pandas as pd

from engine.semantic import SemanticProfiler
from engine.formula_auditor import FormulaAuditor
from engine.impact_simulator import ImpactSimulator
from engine.rules import PlainEnglishRuleEngine
from engine.samples import SampleGenerator
from engine.exporter import SafeSpreadsheetExporter

app = FastAPI(title="Spreadsheet Quality Inspector API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STORAGE_DIR = os.path.join(BASE_DIR, "storage")
SAMPLES_DIR = os.path.join(STORAGE_DIR, "samples")
EXPORTS_DIR = os.path.join(STORAGE_DIR, "exports")
os.makedirs(SAMPLES_DIR, exist_ok=True)
os.makedirs(EXPORTS_DIR, exist_ok=True)

# Generate samples at startup
SAMPLE_FILES = SampleGenerator.create_all_samples(SAMPLES_DIR)

# In-memory session cache for fast interactive reviews
ACTIVE_SESSIONS: Dict[str, Dict[str, Any]] = {}

class EvaluateRuleRequest(BaseModel):
    session_id: str
    rule_text: str

class ExportRequest(BaseModel):
    session_id: str
    accepted_patch_ids: List[str]
    mode: str = "annotated" # "clean" or "annotated"

@app.get("/api/health")
def health():
    return {"status": "ok", "app": "Spreadsheet Quality Inspector", "version": "1.0.0"}

@app.get("/api/samples")
def get_samples():
    return [
        {
            "id": "saas_forecast",
            "name": "SaaS Financial Forecast & Unit Economics (P&L)",
            "description": "Board financial model containing hidden revenue rows, slipped annual sum formulas omitting Q4, and hardcoded formula breaks.",
            "file_type": "xlsx"
        },
        {
            "id": "clinical_intake",
            "name": "Clinical Patient Encounters & Admissions",
            "description": "Hospital intake ledger featuring leading-zero MRNs/ZIPs, ambiguous month/day date conventions, discharge before admission dates, and duplicate encounters.",
            "file_type": "xlsx"
        },
        {
            "id": "ecommerce_orders",
            "name": "E-Commerce Supply Chain & Order Ledger",
            "description": "Order processing workbook with duplicate order numbers, manual formula overrides on line totals, and a sum formula truncating recent transactions.",
            "file_type": "xlsx"
        }
    ]

@app.post("/api/load-sample/{sample_id}")
def load_sample(sample_id: str):
    if sample_id not in SAMPLE_FILES:
        raise HTTPException(status_code=404, detail="Sample not found")

    filepath = SAMPLE_FILES[sample_id]
    session_id = str(uuid.uuid4())
    session_copy = os.path.join(STORAGE_DIR, f"{session_id}.xlsx")
    shutil.copyfile(filepath, session_copy)

    analysis = _run_full_analysis(session_copy, session_id, os.path.basename(filepath))
    ACTIVE_SESSIONS[session_id] = {
        "filepath": session_copy,
        "filename": os.path.basename(filepath),
        "analysis": analysis
    }
    return analysis

@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    filename = file.filename or "uploaded.xlsx"
    ext = os.path.splitext(filename)[1].lower()

    if ext not in [".xlsx", ".csv"]:
        raise HTTPException(status_code=400, detail="Only .xlsx and .csv files are supported.")

    session_id = str(uuid.uuid4())
    save_path = os.path.join(STORAGE_DIR, f"{session_id}{ext}")

    with open(save_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    analysis = _run_full_analysis(save_path, session_id, filename)
    ACTIVE_SESSIONS[session_id] = {
        "filepath": save_path,
        "filename": filename,
        "analysis": analysis
    }
    return analysis

@app.post("/api/evaluate-rule")
def evaluate_rule(req: EvaluateRuleRequest):
    session = ACTIVE_SESSIONS.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session expired or not found")

    analysis = session["analysis"]
    data = analysis["data_preview"]
    columns = analysis["columns"]

    result = PlainEnglishRuleEngine.compile_and_evaluate(req.rule_text, data, columns)
    return result

@app.post("/api/export")
def export_file(req: ExportRequest):
    session = ACTIVE_SESSIONS.get(req.session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session expired or not found")

    source_path = session["filepath"]
    if not source_path.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Export with annotations is currently supported for Excel (.xlsx) files.")

    all_findings = session["analysis"]["findings"]
    accepted_patches = [f for f in all_findings if f["id"] in req.accepted_patch_ids]

    export_filename = f"Cleaned_{session['filename']}"
    export_path = os.path.join(EXPORTS_DIR, f"{uuid.uuid4()}_{export_filename}")

    SafeSpreadsheetExporter.export_cleaned_workbook(
        source_path=source_path,
        output_path=export_path,
        accepted_patches=accepted_patches,
        mode=req.mode
    )

    return {
        "download_url": f"/api/download/{os.path.basename(export_path)}",
        "filename": export_filename,
        "patches_applied": len(accepted_patches)
    }

@app.get("/api/download/{filename}")
def download_file(filename: str):
    path = os.path.join(EXPORTS_DIR, filename)
    if not os.path.exists(path):
        # Check samples dir
        path = os.path.join(SAMPLES_DIR, filename)
        if not os.path.exists(path):
            raise HTTPException(status_code=404, detail="File not found")

    return FileResponse(
        path=path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

def _run_full_analysis(filepath: str, session_id: str, original_filename: str) -> Dict[str, Any]:
    ext = os.path.splitext(filepath)[1].lower()
    data_rows: List[Dict[str, Any]] = []
    columns: List[str] = []
    raw_cells = {}
    formula_audit = {"findings": [], "summary_cells": []}

    if ext == ".xlsx":
        wb_data = openpyxl.load_workbook(filepath, data_only=True)
        wb_formulas = openpyxl.load_workbook(filepath, data_only=False)
        sheet_name = wb_data.sheetnames[0]
        ws_data = wb_data[sheet_name]
        ws_form = wb_formulas[sheet_name]

        # Extract rows
        rows_gen = ws_data.iter_rows(values_only=True)
        header = next(rows_gen, None)
        if header:
            columns = [str(h or f"Column_{i+1}") for i, h in enumerate(header)]
            for r_idx, r in enumerate(rows_gen, start=2):
                if any(v is not None for v in r):
                    row_dict = {}
                    for col_idx, col_name in enumerate(columns):
                        val = r[col_idx] if col_idx < len(r) else None
                        row_dict[col_name] = val
                    data_rows.append(row_dict)

        # Audit formula structures with evaluated data workbook
        formula_audit = FormulaAuditor.audit_workbook(wb_formulas, wb_data)

    else: # CSV
        df = pd.read_csv(filepath, dtype=str)
        columns = list(df.columns)
        data_rows = df.to_dict(orient="records")

    # 1. Semantic Profiling
    semantics = SemanticProfiler.profile_dataset(data_rows, raw_cells)

    # 2. Tabular Integrity Checks (Duplicates, temporal order)
    tabular_findings = _check_tabular_anomalies(data_rows, semantics, columns)

    # 3. Combine findings
    all_findings = []
    all_findings.extend(formula_audit.get("findings", []))
    all_findings.extend(tabular_findings)

    # 4. Impact Simulation (calculate Δ on KPIs)
    enriched_findings = ImpactSimulator.simulate_impacts(
        data=data_rows,
        findings=all_findings,
        profiles=semantics.get("columns", {})
    )

    # 5. Calculate "Can I Trust This Report?" Executive Scorecard
    trust_scorecard = _calculate_trust_score(enriched_findings, semantics, len(data_rows))

    return {
        "session_id": session_id,
        "filename": original_filename,
        "total_rows": len(data_rows),
        "total_columns": len(columns),
        "columns": columns,
        "semantics": semantics,
        "findings": enriched_findings,
        "trust_scorecard": trust_scorecard,
        "data_preview": data_rows[:100]
    }

def _check_tabular_anomalies(data: List[Dict[str, Any]], semantics: Dict[str, Any], columns: List[str]) -> List[Dict[str, Any]]:
    findings = []
    profiles = semantics.get("columns", {})
    candidate_keys = semantics.get("candidate_keys", [])

    # Check key duplicates
    if candidate_keys:
        primary_key = candidate_keys[0]
        seen = {}
        for r_idx, row in enumerate(data, start=2):
            if len(primary_key) == 1:
                key_val = str(row.get(primary_key[0], "")).strip()
            else:
                key_val = "_".join(str(row.get(k, "")).strip() for k in primary_key)

            if not key_val:
                continue

            if key_val in seen:
                findings.append({
                    "id": f"dup_{r_idx}_{primary_key[0]}",
                    "category": "grain_integrity",
                    "severity": "high",
                    "type": "duplicate_key",
                    "cell": f"{primary_key[0]}{r_idx}",
                    "row": r_idx,
                    "column": primary_key[0],
                    "column_name": " + ".join(primary_key),
                    "current_value": key_val,
                    "proposed_value": "Remove or update key",
                    "message": f"Duplicate primary entity key '{key_val}' violates row grain (first seen in row {seen[key_val]}).",
                    "impact_hint": "Causes double-counting in joins, summaries, and volume counts.",
                    "can_autofix": False
                })
            else:
                seen[key_val] = r_idx

    # Check temporal sequencing
    date_cols = [c for c, p in profiles.items() if p["inferred_type"] == "date"]
    start_cols = [c for c in date_cols if any(w in c.lower() for w in ["order", "admit", "start"])]
    end_cols = [c for c in date_cols if any(w in c.lower() for w in ["ship", "deliver", "discharge", "end"])]

    if start_cols and end_cols:
        s_col = start_cols[0]
        e_col = end_cols[0]
        for r_idx, row in enumerate(data, start=2):
            v_s = row.get(s_col)
            v_e = row.get(e_col)
            d_s = PlainEnglishRuleEngine._parse_date(v_s)
            d_e = PlainEnglishRuleEngine._parse_date(v_e)
            if d_s and d_e and d_e < d_s:
                findings.append({
                    "id": f"temporal_{r_idx}_{e_col}",
                    "category": "temporal_logic",
                    "severity": "high",
                    "type": "temporal_sequence_violation",
                    "cell": f"{e_col}{r_idx}",
                    "row": r_idx,
                    "column": e_col,
                    "column_name": e_col,
                    "current_value": str(v_e),
                    "proposed_value": f">={str(v_s)}",
                    "message": f"{e_col} ({v_e}) occurs before {s_col} ({v_s}).",
                    "impact_hint": "Negative turnaround time corrupts operational SLAs.",
                    "can_autofix": False
                })

    # Check date ambiguities
    for amb in semantics.get("date_ambiguities", []):
        findings.append({
            "id": f"date_amb_{amb['column']}",
            "category": "semantic_typing",
            "severity": "medium",
            "type": "date_convention_ambiguity",
            "cell": f"{amb['column']}1",
            "row": 1,
            "column": amb["column"],
            "column_name": amb["column"],
            "current_value": amb.get("sample"),
            "proposed_value": "Explicit ISO YYYY-MM-DD",
            "message": amb.get("message"),
            "impact_hint": "Misinterpreting MM/DD as DD/MM will shift monthly reporting by quarters.",
            "can_autofix": False
        })

    return findings

def _calculate_trust_score(findings: List[Dict[str, Any]], semantics: Dict[str, Any], total_rows: int) -> Dict[str, Any]:
    score = 100
    critical_count = sum(1 for f in findings if f.get("severity") == "critical" or f.get("impact", {}).get("impact_level") == "critical")
    high_count = sum(1 for f in findings if f.get("severity") == "high" or f.get("impact", {}).get("impact_level") == "high")
    medium_count = sum(1 for f in findings if f.get("severity") == "medium")

    score -= (critical_count * 18)
    score -= (high_count * 9)
    score -= (medium_count * 4)
    score = max(min(score, 100), 15)

    # Trust categories
    formula_breaks = sum(1 for f in findings if f.get("category") == "formula_integrity")
    range_slippages = sum(1 for f in findings if f.get("category") == "range_slippage")
    grain_issues = sum(1 for f in findings if f.get("category") == "grain_integrity")
    temporal_issues = sum(1 for f in findings if f.get("category") == "temporal_logic")

    formula_score = max(100 - (formula_breaks * 25 + range_slippages * 30), 20)
    grain_score = max(100 - (grain_issues * 25), 30)
    temporal_score = max(100 - (temporal_issues * 20), 40)

    if score >= 85:
        verdict = "High Confidence - Ready for Decision Making"
        color = "emerald"
    elif score >= 65:
        verdict = "Proceed with Caution - Review Critical Findings"
        color = "amber"
    else:
        verdict = "Do Not Rely - Significant Calculation & Structural Flaws"
        color = "rose"

    return {
        "overall_score": score,
        "verdict": verdict,
        "verdict_color": color,
        "pillars": {
            "formula_uniformity": formula_score,
            "range_coverage": max(100 - (range_slippages * 35), 20),
            "grain_purity": grain_score,
            "temporal_consistency": temporal_score
        },
        "critical_issues": critical_count,
        "high_issues": high_count,
        "medium_issues": medium_count,
        "executive_summary": (
            f"Detected {len(findings)} analytical risks ({critical_count} critical). "
            f"{'Slipped summary formulas and formula pattern disruptions were found.' if formula_breaks or range_slippages else 'No structural formula breaks.'}"
        )
    }
