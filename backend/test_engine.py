import os
import openpyxl
from engine.samples import SampleGenerator
from engine.semantic import SemanticProfiler
from engine.formula_auditor import FormulaAuditor
from engine.impact_simulator import ImpactSimulator
from engine.rules import PlainEnglishRuleEngine
from engine.exporter import SafeSpreadsheetExporter

def test_engine_end_to_end():
    test_dir = os.path.join(os.path.dirname(__file__), "test_output")
    os.makedirs(test_dir, exist_ok=True)

    print("--- 1. Testing Sample Generation ---")
    samples = SampleGenerator.create_all_samples(test_dir)
    assert os.path.exists(samples["saas_forecast"]), "SaaS forecast file should exist"
    assert os.path.exists(samples["clinical_intake"]), "Clinical intake file should exist"
    assert os.path.exists(samples["ecommerce_orders"]), "E-commerce file should exist"
    print("[OK] All 3 sample spreadsheets created successfully.")

    print("\n--- 2. Testing Formula Auditor on SaaS Model ---")
    wb_formulas = openpyxl.load_workbook(samples["saas_forecast"], data_only=False)
    wb_data = openpyxl.load_workbook(samples["saas_forecast"], data_only=True)
    audit = FormulaAuditor.audit_workbook(wb_formulas, wb_data)
    findings = audit["findings"]
    print(f"Detected {len(findings)} formula findings in SaaS model.")

    hardcoded_breaks = [f for f in findings if f["type"] == "hardcoded_override"]
    assert len(hardcoded_breaks) > 0, "Should detect hardcoded override in Net_Margin column"
    print(f"[OK] Detected hardcoded formula break at cell {hardcoded_breaks[0]['cell']}: {hardcoded_breaks[0]['message']}")

    range_slippages = [f for f in findings if f["type"] == "omitted_rows_in_sum"]
    assert len(range_slippages) > 0, "Should detect range slippage in SUM() formula"
    print(f"[OK] Detected range boundary slippage at cell {range_slippages[0]['cell']}: {range_slippages[0]['message']}")

    print("\n--- 3. Testing Semantic Profiling on Clinical Intake ---")
    wb_clin = openpyxl.load_workbook(samples["clinical_intake"], data_only=True)
    ws = wb_clin.active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    data = [dict(zip(headers, r)) for r in rows[1:]]

    semantics = SemanticProfiler.profile_dataset(data)
    mrn_prof = semantics["columns"]["Patient_MRN"]
    zip_prof = semantics["columns"]["Zip_Code"]
    assert mrn_prof["preserve_leading_zeros"] == True, "Patient_MRN must preserve leading zeros"
    assert zip_prof["preserve_leading_zeros"] == True, "Zip_Code must preserve leading zeros"
    print(f"[OK] Leading zeros correctly safeguarded for MRN and Zip Code. Apparent grain: {semantics['grain']}")

    assert len(semantics["date_ambiguities"]) > 0, "Should detect date format ambiguity"
    print(f"[OK] Detected date convention ambiguity: {semantics['date_ambiguities'][0]['message']}")

    print("\n--- 4. Testing Counterfactual Impact Simulator ---")
    enriched_findings = ImpactSimulator.simulate_impacts(data, findings, semantics["columns"])
    for f in enriched_findings[:3]:
        print(f"* Impact [{f['impact']['impact_level'].upper()}]: {f['impact']['executive_statement']}")
    print("[OK] Impact simulator successfully computed counterfactual deltas.")

    print("\n--- 5. Testing Plain English Rule Engine ---")
    rule_res = PlainEnglishRuleEngine.compile_and_evaluate("Patient_MRN should be unique", data, headers)
    print(f"* Plain rule evaluated: '{rule_res['interpreted_as']}' - Passed: {rule_res['passed']}")
    assert rule_res["passed"] == True, "Patient_MRN is unique"

    rule_res_date = PlainEnglishRuleEngine.compile_and_evaluate("Discharge_Date must be after Admission_Date", data, headers)
    print(f"* Plain rule evaluated: '{rule_res_date['interpreted_as']}' - Passed: {rule_res_date['passed']}, Violations: {rule_res_date['violations_count']}")
    assert rule_res_date["passed"] == False, "Should catch Discharge before Admission error"

    print("\n--- 6. Testing Non-Destructive Safe Exporter with Audit Tab ---")
    cleaned_out = os.path.join(test_dir, "cleaned_saas_forecast.xlsx")
    accepted_patches = [hardcoded_breaks[0], range_slippages[0]]
    SafeSpreadsheetExporter.export_cleaned_workbook(
        source_path=samples["saas_forecast"],
        output_path=cleaned_out,
        accepted_patches=accepted_patches,
        mode="annotated"
    )
    assert os.path.exists(cleaned_out), "Exported cleaned workbook should exist"
    wb_exported = openpyxl.load_workbook(cleaned_out)
    assert "_Audit_Changelog" in wb_exported.sheetnames, "Exported workbook must contain _Audit_Changelog sheet"
    print("[OK] Safe non-destructive export generated with green cell highlights and _Audit_Changelog tab.")

    print("\n================ ALL BACKEND ENGINE TESTS PASSED ================")

if __name__ == "__main__":
    test_engine_end_to_end()
